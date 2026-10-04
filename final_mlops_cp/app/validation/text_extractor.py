"""
app/validation/text_extractor.py
==================================
Extracts plain text from PDF, DOCX, and TXT documents.

Preserves document_id and metadata through extraction.
All extracted text is returned as a structured dict ready for chunking.
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


class TextExtractor:
    """
    Extracts text from supported file formats.

    Returns a dict:
        {
            "document_id": str,
            "file_path":   str,
            "text":        str,
            "metadata":    dict,
            "page_count":  int,   # PDFs only; 1 for others
            "char_count":  int,
        }
    """

    def extract(self, file_path: Path, metadata: Optional[dict] = None) -> dict:
        """Extract text from a single file."""
        ext = file_path.suffix.lower()

        if ext == ".pdf":
            text, pages = self._extract_pdf(file_path)
        elif ext == ".docx":
            text, pages = self._extract_docx(file_path)
        elif ext == ".txt":
            text, pages = self._extract_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

        return {
            "document_id": (metadata or {}).get("document_id", file_path.stem),
            "file_path": str(file_path),
            "text": text,
            "metadata": metadata or {},
            "page_count": pages,
            "char_count": len(text),
        }

    def extract_batch(
        self, file_paths: list[Path], metadata_dir: Optional[Path] = None
    ) -> list[dict]:
        """Extract text from a list of files."""
        results = []
        for fp in file_paths:
            meta = self._load_metadata(fp, metadata_dir) if metadata_dir else None
            try:
                result = self.extract(fp, metadata=meta)
                if result["char_count"] == 0:
                    logger.warning(f"Empty text extracted from: {fp}")
                else:
                    results.append(result)
                    logger.debug(
                        f"Extracted {result['char_count']} chars from "
                        f"{result['document_id']} ({fp.name})"
                    )
            except Exception as exc:
                logger.error(f"Extraction failed for {fp}: {exc}")
        return results

    # -----------------------------------------------------------------------
    # Format-specific extractors
    # -----------------------------------------------------------------------

    def _extract_pdf(self, file_path: Path) -> tuple[str, int]:
        """Extract text from PDF using PyMuPDF (fitz)."""
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(str(file_path))
            pages = doc.page_count
            texts = []
            for page in doc:
                page_text = page.get_text("text")
                if page_text.strip():
                    texts.append(page_text)
            doc.close()
            return "\n\n".join(texts), pages
        except ImportError:
            raise RuntimeError(
                "PyMuPDF not installed. Run: pip install pymupdf"
            )

    def _extract_docx(self, file_path: Path) -> tuple[str, int]:
        """Extract text from DOCX using python-docx."""
        try:
            from docx import Document
            doc = Document(str(file_path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        paragraphs.append(row_text)
            return "\n\n".join(paragraphs), 1
        except ImportError:
            raise RuntimeError(
                "python-docx not installed. Run: pip install python-docx"
            )

    def _extract_txt(self, file_path: Path) -> tuple[str, int]:
        """Extract text from plain text file with encoding detection."""
        # Try UTF-8 first, then detect encoding
        try:
            text = file_path.read_text(encoding="utf-8")
            return text, 1
        except UnicodeDecodeError:
            try:
                import chardet
                raw = file_path.read_bytes()
                detected = chardet.detect(raw)
                encoding = detected.get("encoding", "latin-1") or "latin-1"
                text = raw.decode(encoding, errors="replace")
                return text, 1
            except ImportError:
                text = file_path.read_bytes().decode("latin-1", errors="replace")
                return text, 1

    def _load_metadata(self, file_path: Path, metadata_dir: Path) -> Optional[dict]:
        """Load the metadata JSON sidecar for a given file."""
        if not metadata_dir.exists():
            return None
        for json_file in metadata_dir.glob("*.json"):
            if json_file.name == "document_index.json":
                continue
            try:
                data = json.loads(json_file.read_text(encoding="utf-8"))
                stored = data.get("file_path", "").replace("\\", "/")
                needle = str(file_path).replace("\\", "/")
                if file_path.name in stored or needle.endswith(stored.lstrip("./")):
                    return data
            except Exception:
                continue
        return None
